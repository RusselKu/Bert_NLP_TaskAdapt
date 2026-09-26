from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "final_report.md"
OUT_PATH = ROOT / "docs" / "final_report.docx"

LOGO_UPY = ROOT / "docs" / "assets" / "logo_upy.png"
LOGO_RENACIMIENTO = ROOT / "docs" / "assets" / "logo_renacimiento.png"

TEAM = [
    "Damian Nicolas Sanchez Novelo",
    "Russel Emmanuel Ku Aguilar",
    "Angel Rivaldo Canche Chuc",
    "Bianca Alexandra Acosta Castellanos",
    "Jonathan Abisai Velasco Martin",
]


def set_cell_shading(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def add_inline_markdown(paragraph, text):
    # Minimal inline Markdown support:
    # **bold**, `code`, *italic*, and ordinary text.
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    pos = 0

    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])

        token = match.group(0)

        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True

        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)

        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True

        pos = match.end()

    if pos < len(text):
        paragraph.add_run(text[pos:])


def clean_markdown(text):
    # Keep content intact while avoiding glyph issues in some PDF renderers.
    return (
        text.replace("—", "-")
            .replace("–", "-")
            .replace("≈", "approximately ")
    )


def add_markdown_table(doc, lines):
    rows = []

    for line in lines:
        row = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(row)

    # Remove Markdown separator row.
    if len(rows) >= 2:
        separator = rows[1]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in separator):
            rows.pop(1)

    if not rows:
        return

    cols = max(len(r) for r in rows)

    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            value = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            add_inline_markdown(p, value)

            for run in p.runs:
                run.font.size = Pt(8.5)

            if r_idx == 0:
                set_cell_shading(cell)
                for run in p.runs:
                    run.bold = True

    set_repeat_table_header(table.rows[0])

    doc.add_paragraph()


def build():
    raw = MD_PATH.read_text(encoding="utf-8")
    raw = clean_markdown(raw)
    lines = raw.splitlines()

    doc = Document()

    # ------------------------------------------------------------
    # PAGE LAYOUT
    # ------------------------------------------------------------
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # ------------------------------------------------------------
    # GLOBAL STYLES
    # ------------------------------------------------------------
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in [
        ("Title", 24),
        ("Heading 1", 16),
        ("Heading 2", 13),
        ("Heading 3", 11),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True

    # ------------------------------------------------------------
    # COVER
    # ------------------------------------------------------------

    # Institutional logos
    logo_table = doc.add_table(rows=1, cols=2)
    logo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    logo_table.autofit = False

    left = logo_table.cell(0, 0)
    right = logo_table.cell(0, 1)

    left.width = Cm(8.0)
    right.width = Cm(8.0)

    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p_logo = left.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_UPY.exists():
        p_logo.add_run().add_picture(
            str(LOGO_UPY),
            width=Inches(2.35)
        )

    p_logo = right.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_RENACIMIENTO.exists():
        p_logo.add_run().add_picture(
            str(LOGO_RENACIMIENTO),
            width=Inches(2.25)
        )

    # Remove logo-table borders
    tblPr = logo_table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)

    tblPr.append(borders)

    # Main academic information
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("U2T01")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(13)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Adapting BERT for NLP Tasks")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(22)

    cover_lines = [
        ("Universidad Politécnica de Yucatán", True),
        ("Data Engineering", True),
        ("Trends in Data Science", True),
        ("Group B", True),
    ]

    for value, italic in cover_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(value)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.italic = italic

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Instructor: Dexter Gomez")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Team Members")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)

    for member in TEAM:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(member)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("26 September 2026")
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    doc.add_page_break()

    # ------------------------------------------------------------
    # HEADER / FOOTER
    # ------------------------------------------------------------
    header = section.header.paragraphs[0]
    header.text = "U2T01 - Adapting BERT for NLP Tasks"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for run in header.runs:
        run.font.size = Pt(8)
        run.font.name = "Arial"

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    # ------------------------------------------------------------
    # CONTENT PARSER
    # ------------------------------------------------------------
    i = 0
    first_h1_skipped = False

    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # Skip Markdown title because cover already contains it.
        if line.startswith("# ") and not first_h1_skipped:
            first_h1_skipped = True
            i += 1
            continue

        # Images
        img_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if img_match:
            alt, rel_path = img_match.groups()
            image_path = (MD_PATH.parent / rel_path).resolve()

            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()

                try:
                    run.add_picture(str(image_path), width=Inches(6.05))
                except Exception as exc:
                    p.add_run(f"[Image could not be inserted: {image_path.name}]")
                    print(f"WARNING image {image_path}: {exc}")
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(f"[Missing image: {rel_path}]")
                print(f"WARNING missing image: {image_path}")

            i += 1
            continue

        # Markdown tables
        if line.lstrip().startswith("|"):
            table_lines = []

            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1

            add_markdown_table(doc, table_lines)
            continue

        # Headings
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        # Bullet items
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, line[2:].strip())
            i += 1
            continue

        # Numbered items
        if re.match(r"^\d+\.\s+", line):
            value = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, value)
            i += 1
            continue

        # Caption: markdown italic paragraph "*Figure ...*"
        if line.startswith("*Figure ") and line.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[1:-1])
            r.italic = True
            r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        # Horizontal rule
        if line.strip() in {"---", "***"}:
            i += 1
            continue

        # Ordinary paragraph.
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_markdown(p, line.strip())

        i += 1

    # Keep headings attached to following paragraphs where possible.
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(5)

    # Core document properties
    props = doc.core_properties
    props.title = "U2T01 - Adapting BERT for NLP Tasks"
    props.subject = "Final Experimental Report"
    props.keywords = "BERT, NLP, Fine-Tuning, NER, POS, QA, Topic Classification"

    doc.save(OUT_PATH)

    print("=" * 72)
    print("DOCX GENERATED")
    print("=" * 72)
    print(OUT_PATH)
    print(f"Size: {OUT_PATH.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    build()
